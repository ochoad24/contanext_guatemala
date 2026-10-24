import csv
import hashlib
import io
import zipfile
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation


TEMPLATE_VERSION = "GT-PHYSICAL-COUNT-1"
REQUIRED_COLUMNS = ("item_code", "warehouse", "uom", "counted_quantity")
OPTIONAL_COLUMNS = ("note",)
MAX_FILE_BYTES = 25 * 1024 * 1024
MAX_ROWS = 100_000


@dataclass(frozen=True, slots=True)
class ParsedCountRow:
	source_row: int
	item_code: str
	warehouse: str
	uom: str
	counted_quantity: Decimal | None
	note: str
	status: str
	error: str = ""


def sha256_hex(content: bytes) -> str:
	return hashlib.sha256(content).hexdigest()


def parse_physical_count(content: bytes, filename: str) -> list[ParsedCountRow]:
	if not content:
		raise ValueError("El archivo de conteo está vacío.")
	if len(content) > MAX_FILE_BYTES:
		raise ValueError("El archivo de conteo excede 25 MiB.")
	filename = (filename or "").lower()
	if filename.endswith(".csv"):
		rows = _read_csv(content)
	elif filename.endswith(".xlsx"):
		rows = _read_xlsx(content)
	else:
		raise ValueError("El conteo físico debe cargarse en CSV o XLSX.")
	if len(rows) > MAX_ROWS:
		raise ValueError(f"El conteo excede el máximo técnico de {MAX_ROWS} filas.")
	return [_normalize_row(index, row) for index, row in enumerate(rows, start=2)]


def _read_csv(content: bytes) -> list[dict]:
	try:
		text = content.decode("utf-8-sig")
	except UnicodeDecodeError as error:
		raise ValueError("El CSV debe usar codificación UTF-8.") from error
	try:
		reader = csv.DictReader(io.StringIO(text), strict=True)
		_validate_headers(reader.fieldnames or [])
		rows = []
		for row in reader:
			rows.append(dict(row))
			if len(rows) > MAX_ROWS:
				raise ValueError(f"El conteo excede el máximo técnico de {MAX_ROWS} filas.")
		return rows
	except csv.Error as error:
		raise ValueError("El CSV no tiene una estructura válida.") from error


def _read_xlsx(content: bytes) -> list[dict]:
	from openpyxl import load_workbook

	_validate_xlsx_container(content)
	try:
		workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
	except Exception as error:
		raise ValueError("El XLSX no es válido o no se puede leer.") from error
	worksheet = workbook.active
	iterator = worksheet.iter_rows(values_only=False)
	try:
		header_cells = next(iterator)
	except StopIteration:
		raise ValueError("El XLSX no contiene encabezados.")
	headers = [str(cell.value or "").strip() for cell in header_cells]
	_validate_headers(headers)
	rows = []
	for cells in iterator:
		if any(cell.data_type == "f" for cell in cells):
			raise ValueError("El XLSX no puede contener fórmulas.")
		values = [cell.value for cell in cells]
		if not any(value not in (None, "") for value in values):
			continue
		rows.append(dict(zip(headers, values, strict=False)))
		if len(rows) > MAX_ROWS:
			raise ValueError(f"El conteo excede el máximo técnico de {MAX_ROWS} filas.")
	workbook.close()
	return rows


def _validate_xlsx_container(content: bytes) -> None:
	try:
		with zipfile.ZipFile(io.BytesIO(content)) as archive:
			members = archive.infolist()
			if any(member.flag_bits & 0x1 for member in members):
				raise ValueError("El XLSX no puede estar cifrado.")
			if any(".." in member.filename.split("/") for member in members):
				raise ValueError("El XLSX contiene rutas no permitidas.")
			if sum(member.file_size for member in members) > 100 * 1024 * 1024:
				raise ValueError("El contenido expandido del XLSX excede 100 MiB.")
	except zipfile.BadZipFile as error:
		raise ValueError("El XLSX no es un contenedor válido.") from error


def _validate_headers(headers: list[str]) -> None:
	normalized = [str(value or "").strip() for value in headers]
	if len(normalized) != len(set(normalized)):
		raise ValueError("La plantilla contiene columnas duplicadas.")
	missing = [column for column in REQUIRED_COLUMNS if column not in normalized]
	unknown = [column for column in normalized if column not in REQUIRED_COLUMNS + OPTIONAL_COLUMNS]
	if missing or unknown:
		details = []
		if missing:
			details.append("faltan: " + ", ".join(missing))
		if unknown:
			details.append("desconocidas: " + ", ".join(unknown))
		raise ValueError("Columnas inválidas (" + "; ".join(details) + ").")


def _normalize_row(source_row: int, source: dict) -> ParsedCountRow:
	item_code = str(source.get("item_code") or "").strip()
	warehouse = str(source.get("warehouse") or "").strip()
	uom = str(source.get("uom") or "").strip()
	note = str(source.get("note") or "").strip()
	errors = []
	for value, label in ((item_code, "item_code"), (warehouse, "warehouse"), (uom, "uom")):
		if not value:
			errors.append(f"{label} es obligatorio")
		elif value[0] in "=+@":
			errors.append(f"{label} contiene un prefijo no permitido")
	try:
		quantity = Decimal(str(source.get("counted_quantity", "")).strip())
		if not quantity.is_finite():
			raise InvalidOperation
	except (InvalidOperation, ValueError):
		quantity = None
		errors.append("counted_quantity no es decimal")
	return ParsedCountRow(
		source_row=source_row,
		item_code=item_code,
		warehouse=warehouse,
		uom=uom,
		counted_quantity=quantity,
		note=note,
		status="Rejected" if errors else "Accepted",
		error="; ".join(errors),
	)
