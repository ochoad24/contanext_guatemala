import csv
import io
from collections.abc import Iterable, Mapping, Sequence
from datetime import date, datetime
from decimal import Decimal


DANGEROUS_SPREADSHEET_PREFIXES = ("=", "+", "-", "@", "\t", "\r")


def csv_bytes(columns: Sequence[str], rows: Iterable[Mapping]) -> bytes:
	output = io.StringIO(newline="")
	writer = csv.writer(output, dialect="excel", lineterminator="\r\n")
	writer.writerow(columns)
	for row in rows:
		writer.writerow([_export_value(row.get(column)) for column in columns])
	return output.getvalue().encode("utf-8-sig")


def xlsx_bytes(
	columns: Sequence[str], rows: Iterable[Mapping], metadata: Mapping[str, object]
) -> bytes:
	from openpyxl import Workbook

	workbook = Workbook(write_only=True)
	meta_sheet = workbook.create_sheet("Metadatos")
	meta_sheet.append(["clave", "valor"])
	for key in sorted(metadata):
		meta_sheet.append([_export_value(key), _export_value(metadata[key])])
	data_sheet = workbook.create_sheet("Datos")
	data_sheet.append(list(columns))
	for row in rows:
		data_sheet.append([_export_value(row.get(column)) for column in columns])
	content = io.BytesIO()
	workbook.save(content)
	return content.getvalue()


def _export_value(value):
	if value is None:
		return ""
	if isinstance(value, Decimal):
		return str(value)
	if isinstance(value, (date, datetime)):
		return value.isoformat()
	if isinstance(value, bool):
		return "1" if value else "0"
	if isinstance(value, str) and value.startswith(DANGEROUS_SPREADSHEET_PREFIXES):
		return "'" + value
	return value
